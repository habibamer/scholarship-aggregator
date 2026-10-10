import json
import os
import time
from datetime import datetime, timezone

import cloudscraper
import openpyxl
import requests
from bs4 import BeautifulSoup
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

EXCEL_FILE = "scholarships.xlsx"
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36"
    ),
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
    "Accept-Language": "en-US,en;q=0.9",
}

scraper = cloudscraper.create_scraper(
    browser={"browser": "chrome", "platform": "windows", "desktop": True}
)
plain = requests.Session()
plain.headers.update(HEADERS)

SOURCES = {"for9a": "for9a.com", "od": "opportunitydesk.org"}

CLOSED_WORDS = (
    "closed", "expired", "مغلق", "منتهي", 
    "closes today", "closed today", 
    "deadline passed", "applications closed",
    "2025",
    "january 2026", "february 2026", "march 2026",
    "jan 2026", "feb 2026", "mar 2026",
    "valentine season – february", "february 10, 2026"
)

def is_closed(r):
    text = f"{r[1]} {r[3]} {r[4]}".lower()
    return any(w in text for w in CLOSED_WORDS)


def source_of(link):
    for name, domain in SOURCES.items():
        if domain in (link or ""):
            return name
    return "other"


def get_html(url, timeout=25, tries=2):
    for client_name, client in (("cloudscraper", scraper), ("requests", plain)):
        for attempt in range(1, tries + 1):
            try:
                res = client.get(url, timeout=timeout)
                if res.status_code == 200 and res.content:
                    return res.content
                if res.status_code in (403, 404):
                    break
            except Exception:
                pass
            time.sleep(2)
    return None


old_rows = []
if os.path.exists(EXCEL_FILE):
    try:
        wb_old = openpyxl.load_workbook(EXCEL_FILE)
        for r in wb_old.active.iter_rows(min_row=2, values_only=True):
            if r and r[1] and r[4]:
                row_str = [str(x) if x is not None else "" for x in r[:5]]
                if not is_closed(row_str):
                    old_rows.append(row_str)
    except Exception as e:
        print(f"[Warn] Could not read old file: {e}")
print(f"[Old data] {len(old_rows)} active rows loaded (closed & old dates excluded).")


def scrape_for9a():
    urls = {
        "bachelor": "https://www.for9a.com/en/opportunity/category/Scholarships/bachelor",
        "master": "https://www.for9a.com/en/opportunity/category/Scholarships/master",
        "phd": "https://www.for9a.com/en/opportunity/category/Scholarships/phd",
    }
    out = []
    for deg, base_url in urls.items():
        print(f"[For9a] Fetching level: {deg}...")
        fails = 0
        for page in range(1, 9):
            content = get_html(f"{base_url}?page={page}")
            if content is None:
                fails += 1
                if fails >= 2:
                    break
                continue
            fails = 0
            soup = BeautifulSoup(content, "lxml")
            cards = soup.find_all("div", class_="p-4 flex flex-col flex-grow")
            if not cards:
                break
            for card in cards:
                text = card.get_text().lower()
                if any(w in text for w in CLOSED_WORDS):
                    continue
                a = card.find("a", class_="editor_page")
                if not a:
                    continue
                href = a.get("href", "")
                link = href if href.startswith("http") else f"https://www.for9a.com{href}"
                c_tag = card.find("span", class_="bg-gray-100")
                country = c_tag.get_text(strip=True) if c_tag else "International"
                d_tag = card.find("span", class_=lambda x: x and "bg-orange-50" in x)
                deadline = d_tag.get_text(strip=True) if d_tag else "N/A"
                item = [deg, a.get_text(strip=True), country, deadline, link]
                if not is_closed(item) and item not in out:
                    out.append(item)
            time.sleep(1)
    return out


def parse_od_html(content):
    soup = BeautifulSoup(content, "lxml")
    articles = soup.find_all("article")
    if not articles:
        articles = soup.find_all("div", class_=lambda x: x and "post" in x)
    items = []
    for art in articles:
        full_art_html = str(art).lower()
        if any(w in full_art_html for w in CLOSED_WORDS):
            continue
            
        t = art.find(["h2", "h3"])
        a = t.find("a") if t else None
        if a and a.get("href"):
            title_text = a.get_text(strip=True)
            link_val = a["href"]
            combined = f"{title_text} {full_art_html}".lower()
            if not any(w in combined for w in CLOSED_WORDS):
                items.append((title_text, link_val))
    return items


def parse_od_feed(content):
    soup = BeautifulSoup(content, "xml")
    items = []
    for it in soup.find_all("item"):
        t = it.find("title")
        l = it.find("link")
        desc = it.find("description")
        content_encoded = it.find("content:encoded")
        
        full_text = ""
        if t: full_text += t.get_text(strip=True) + " "
        if desc: full_text += desc.get_text(strip=True) + " "
        if content_encoded: full_text += content_encoded.get_text(strip=True) + " "
        
        full_text_lower = full_text.lower()
        
        if any(w in full_text_lower for w in CLOSED_WORDS):
            continue
            
        if t and l and l.get_text(strip=True):
            title_text = t.get_text(strip=True)
            link_text = l.get_text(strip=True)
            if not any(w in title_text.lower() for w in CLOSED_WORDS):
                items.append((title_text, link_text))
    return items


def scrape_od():
    urls = {
        "bachelor": "https://opportunitydesk.org/category/fellowships-and-scholarships/undergraduate/",
        "master": "https://opportunitydesk.org/category/fellowships-and-scholarships/masters-postgraduate/",
        "phd": "https://opportunitydesk.org/category/fellowships-and-scholarships/phd-post-doctoral/",
    }
    out = []
    for deg, base in urls.items():
        print(f"\n[OpportunityDesk] Fetching level: {deg}...")
        fails = 0
        for page in range(1, 6):
            page_url = base if page == 1 else f"{base}page/{page}/"
            feed_url = base + "feed/" + ("" if page == 1 else f"?paged={page}")

            items = []
            html = get_html(page_url)
            if html:
                items = parse_od_html(html)
            if not items:
                feed = get_html(feed_url)
                if feed:
                    items = parse_od_feed(feed)

            if not items:
                fails += 1
                if fails >= 2:
                    break
                continue
            fails = 0
            for title, link in items:
                item = [deg, title, "International", "Check Official Website", link]
                if not is_closed(item) and item not in out:
                    out.append(item)
            time.sleep(1)
    return out


new_for9a = scrape_for9a()
new_od = scrape_od()
print(f"\n[Result] For9a: {len(new_for9a)} | OpportunityDesk: {len(new_od)}")

final = []
seen = set()


def add(rows):
    for r in rows:
        if not is_closed(r):
            key = (r[0], r[4])
            if key not in seen:
                seen.add(key)
                final.append(r)


for name, new_rows in (("for9a", new_for9a), ("od", new_od)):
    old_src = [r for r in old_rows if source_of(r[4]) == name]
    if len(new_rows) >= 0.5 * len(old_src) and new_rows:
        add(new_rows)
    else:
        add(new_rows)
        add(old_src)
add([r for r in old_rows if source_of(r[4]) == "other"])

if not final:
    raise SystemExit("[ABORT] No data at all.")

before = len(final)
final = [r for r in final if not is_closed(r)]
print(f"[Closed Filter] Removed {before - len(final)} closed or old (up to March 2026) scholarships.")

order = {"bachelor": 1, "master": 2, "phd": 3}
final.sort(key=lambda r: order.get(r[0].lower(), 4))
print(f"[Summary] Total active scholarships in file: {len(final)}")

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Scholarships"
ws.append(["Degree", "Title", "Country", "Deadline", "Link"])

fill = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")
for i in range(1, 6):
    c = ws.cell(row=1, column=i)
    c.fill = fill
    c.font = Font(bold=True, size=11)
    c.alignment = Alignment(horizontal="center", vertical="center")

for row in final:
    ws.append(row)

align = Alignment(wrap_text=True, vertical="center")
for row in ws.iter_rows(min_row=2, max_row=len(final) + 1, min_col=1, max_col=5):
    for c in row:
        c.alignment = align

for idx, w in {1: 15, 2: 45, 3: 18, 4: 25, 5: 40}.items():
    ws.column_dimensions[get_column_letter(idx)].width = w

try:
    wb.save(EXCEL_FILE)
    print(f"Successfully updated {EXCEL_FILE} with {len(final)} items.")
except PermissionError:
    wb.save("scholarships_new.xlsx")

with open("last_updated.json", "w") as f:
    json.dump({"updated": datetime.now(timezone.utc).isoformat()}, f)
